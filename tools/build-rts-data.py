#!/usr/bin/env python3
"""Extract the Risk Theory Society tracking workbook into rts.qmd.

The source of truth is a spreadsheet Evan maintains and edits in place:

    files/RTSPublicationTracking_*.xlsx    (newest by filename wins)

It has two sheets. `Papers` is one row per paper presented at an RTS meeting,
with up to six author/affiliation pairs and, where the paper reached print, a
journal, a publication year, and a link. `Meetings` is one row per meeting with
the president, the host institution, the location, and the society's own member
and attendee counts.

This script normalizes both sheets and writes them as a compact JSON blob into
the marked block inside `rts.qmd`. The page computes every displayed statistic
from that blob at load time, so no number on the page can drift from the
workbook.

The workbook itself is deliberately NOT committed. `files/` is copied to the
published site, so a tracked .xlsx would sit at a guessable public URL while the
page carrying it is unlisted. Keep the workbook local; commit only what this
script emits.

Refresh after editing the workbook:

    py -3 tools/build-rts-data.py

Records are emitted as positional arrays rather than objects. It roughly halves
the payload, and the page has one decoder in one place.

    paper   [year, title, [[author, affil], ...], journal, year_published, link]
    meeting [year, president, president_affil, host, location, date,
             paper_count, members, attendees]
"""

import json
import re
import sys
from datetime import date, datetime
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parent.parent
WORKBOOK_GLOB = "RTSPublicationTracking_*.xlsx"
PAGE = ROOT / "rts.qmd"


def find_workbook():
    """Pick the newest tracking workbook in files/.

    The filename carries a year, so a refresh may arrive as a new file rather
    than an edit to the old one. Sorting by name puts the highest year last;
    ignoring Excel's `~$` lock files keeps an open workbook from being chosen.
    """
    candidates = sorted(
        path for path in (ROOT / "files").glob(WORKBOOK_GLOB)
        if not path.name.startswith("~$")
    )
    return candidates[-1] if candidates else None

BEGIN = "<!-- BEGIN GENERATED DATA -->"
END = "<!-- END GENERATED DATA -->"

# The workbook uses "." as its empty marker, not a blank cell. Treating "." as a
# value is how a placeholder becomes a journal named "." in a chart.
EMPTY = {"", ".", "..", "n/a", "N/A", "na", "NA", "-", "--", "none", "None"}

warnings = []


def warn(message):
    warnings.append(message)


def text(value):
    """Normalize a cell to a clean string, or "" for the sheet's empty markers."""
    if value is None:
        return ""
    if isinstance(value, (datetime, date)):
        # Excel coerced a text date range into a real date. The day-of-month may
        # survive but the year is the coercion's, not the meeting's, so this is
        # reported rather than rendered.
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    cleaned = re.sub(r"\s+", " ", str(value)).strip()
    return "" if cleaned in EMPTY else cleaned


def integer(value):
    """Normalize a cell to an int, or None when absent or non-numeric."""
    raw = text(value)
    if not raw:
        return None
    try:
        return int(float(raw))
    except ValueError:
        return None


def read_papers(sheet):
    papers = []
    for index, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
        year = integer(row[0])
        title = text(row[1])
        if year is None and not title:
            continue
        if year is None:
            warn(f"Papers row {index}: no meeting year, row skipped ({title[:50]!r})")
            continue
        if not title:
            warn(f"Papers row {index}: no title, row skipped ({year})")
            continue

        # Author/affiliation pairs occupy columns 2..13. An author with a blank
        # affiliation is kept; an affiliation with no author is not a person.
        authors = []
        for column in range(2, 14, 2):
            name = text(row[column])
            if name:
                authors.append([name, text(row[column + 1])])
        if not authors:
            warn(f"Papers row {index}: no authors ({year}, {title[:40]!r})")

        journal = text(row[14])
        published = integer(row[15])
        link = text(row[16])
        if link and not link.startswith("http"):
            warn(f"Papers row {index}: link is not a URL, dropped ({link[:40]!r})")
            link = ""

        # A publication year with no journal, or vice versa, is a half-filled
        # row. Report it; the page renders whichever half exists.
        if journal and published is None:
            warn(f"Papers row {index}: journal but no year ({journal})")
        if published is not None and not journal:
            warn(f"Papers row {index}: publication year but no journal ({published})")

        papers.append([year, title, authors, journal, published, link])

    papers.sort(key=lambda p: (p[0], p[1].lower()))
    return papers


def read_meetings(sheet):
    meetings = []
    for index, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
        year = integer(row[0])
        if year is None:
            continue
        if isinstance(row[5], (datetime, date)):
            warn(
                f"Meetings row {index} ({year}): date cell is a real date "
                f"({row[5]:%Y-%m-%d}), not text - Excel coerced a range like "
                f"'April 9-11'. Emitted blank; fix as text in the workbook."
            )
        meetings.append([
            year,
            text(row[1]),
            text(row[2]),
            text(row[3]),
            text(row[4]),
            text(row[5]),
            integer(row[6]),
            integer(row[7]),
            integer(row[8]),
        ])
    meetings.sort(key=lambda m: m[0])
    return meetings


def main():
    workbook = find_workbook()
    if workbook is None:
        print(
            f"build-rts-data: no {WORKBOOK_GLOB} in {ROOT / 'files'}",
            file=sys.stderr,
        )
        return 1
    if not PAGE.exists():
        print(f"build-rts-data: page not found at {PAGE}", file=sys.stderr)
        return 1

    print(f"build-rts-data: reading {workbook.name}")
    book = openpyxl.load_workbook(workbook, data_only=True)
    papers = read_papers(book["Papers"])
    meetings = read_meetings(book["Meetings"])

    payload = {
        "source": workbook.name,
        "papers": papers,
        "meetings": meetings,
    }
    blob = json.dumps(payload, ensure_ascii=False, separators=(",", ":"))

    # </script> inside JSON string data would close the host <script> element
    # early. Escaping the slash keeps the JSON valid and the element intact.
    blob = blob.replace("</", "<\\/")

    block = (
        f'{BEGIN}\n'
        f'<script id="rts-data" type="application/json">\n'
        f'{blob}\n'
        f'</script>\n'
        f'{END}'
    )

    page = PAGE.read_text(encoding="utf-8")
    pattern = re.compile(
        re.escape(BEGIN) + r".*?" + re.escape(END), re.DOTALL
    )
    if not pattern.search(page):
        print(
            f"build-rts-data: no {BEGIN} ... {END} block in {PAGE.name}",
            file=sys.stderr,
        )
        return 1
    PAGE.write_text(pattern.sub(lambda _: block, page, count=1), encoding="utf-8")

    placed = sum(1 for p in papers if p[3])
    years = sorted({p[0] for p in papers})
    print(f"build-rts-data: {len(papers)} papers, {len(meetings)} meetings")
    print(f"build-rts-data: {placed} papers with a journal recorded "
          f"({100 * placed / len(papers):.0f}%)")
    print(f"build-rts-data: meeting years {years[0]}-{years[-1]}, "
          f"{len(years)} present")
    missing = [y for y in range(years[0], years[-1] + 1) if y not in years]
    if missing:
        print(f"build-rts-data: no papers for {compress(missing)}")
    print(f"build-rts-data: wrote {len(blob) / 1024:.0f} KB into {PAGE.name}")
    for message in warnings:
        print(f"build-rts-data: WARNING {message}")
    return 0


def compress(years):
    """Render a sorted year list as ranges: [1985..1993, 2020] -> '1985-1993, 2020'."""
    spans = []
    for year in years:
        if spans and year == spans[-1][1] + 1:
            spans[-1][1] = year
        else:
            spans.append([year, year])
    return ", ".join(
        str(lo) if lo == hi else f"{lo}-{hi}" for lo, hi in spans
    )


if __name__ == "__main__":
    sys.exit(main())
